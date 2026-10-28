from __future__ import annotations

import csv
import itertools
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import time
import unicodedata
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
VALUES_PATH = ROOT / "FrontEnd" / "js" / "data" / "indicator-values.js"
CATALOG_PATH = ROOT / "FrontEnd" / "js" / "data" / "catalog-data.js"
PORDATA_SOURCES_PATH = ROOT / "data" / "pordata-sources.json"
PORDATA_DIRECTORY = ROOT / ".cache" / "pordata"
REPORT_PATH = ROOT / "data" / "update-report.json"
POWERBI_PATH = ROOT / "powerbi-data.csv"
POWERBI_HISTORY_PATH = ROOT / "powerbi-historico.csv"
HISTORY_YEARS = 1000  # nº máximo de anos guardados por indicador (1000 = todos os anos disponíveis)

TERRITORIES = [
    "Alto Minho",
    "Arcos de Valdevez",
    "Caminha",
    "Melgaço",
    "Monção",
    "Paredes de Coura",
    "Ponte da Barca",
    "Ponte de Lima",
    "Valença",
    "Viana do Castelo",
    "Vila Nova de Cerveira",
]

# Estabelecimentos por nível de ensino (Pordata 213): o ficheiro tem uma coluna por nível, em cada ano
# (0 = pré-escolar, 1 = 1.º ciclo, 2 = 2.º ciclo, 3 = 3.º ciclo, 4 = secundário). O indicador é a SOMA dos 5 níveis,
# com posições fixas (sem procurar a combinação mais parecida). Segundo a Pordata, cada estabelecimento é contado
# tantas vezes quantos os ensinos que ministra.
PORDATA_FIXED_POSITIONS = {
    "Estabelecimentos nos ensinos pré-escolar, básico e secundário: por nível de ensino": (0, 1, 2, 3, 4),
}

# Dados base de cada indicador de posição fixa, usados para criar a entrada em indicator-values.js (e em
# catalog-data.js) quando ainda não existe. "catalog" = domínio/tema/categoria no catálogo.
_PF = (0,)  # 1.ª coluna de cada ano = "Total" nos ficheiros "total e por sexo / grupo etário" da Pordata
_ENSINO = "Estabelecimentos nos ensinos pré-escolar, básico e secundário: por nível de ensino"
_RSI = "Beneficiários do Rendimento Mínimo Garantido e Rendimento Social de Inserção da Segurança Social: total e por grupo etário"
_ESTR = "População estrangeira com estatuto legal de residente: total e por sexo"
_SEXO = "População residente (estimativas a 31 de dezembro): total e por sexo"
_GANHO = "Ganho médio mensal dos trabalhadores por conta de outrem: total e por sexo"
PORDATA_FIXED_POSITIONS.update({
    "Freguesias": _PF,
    "Índice de dependência total": _PF,
    _SEXO: _PF,
    _GANHO: _PF,
    _RSI: _PF,
    _ESTR: _PF,
    "Saldo Total": _PF,
})
PORDATA_FIXED_ENTRIES = {
    _ENSINO: {"unit": "estabelecimentos", "dimension": "Soma dos 5 níveis de ensino (cada estabelecimento conta uma vez por cada ensino que ministra)", "prefix": "", "suffix": ""},
    "Freguesias": {"unit": "freguesias", "dimension": "Número de freguesias", "prefix": "", "suffix": "",
                   "catalog": ("Contexto Geográfico", "", "Divisão administrativa")},
    "Índice de dependência total": {"unit": "índice", "dimension": "Jovens e idosos por 100 pessoas em idade ativa", "prefix": "", "suffix": "",
                                    "catalog": ("Demografia", "Dinâmica populacional", "Estrutura da População")},
    _SEXO: {"unit": "pessoas", "dimension": "Total (estimativa a 31 de dezembro)", "prefix": "", "suffix": "",
            "catalog": ("Demografia", "Dinâmica populacional", "Distribuição da população")},
    _GANHO: {"unit": "€", "dimension": "Ganho médio mensal — total", "prefix": "€ ", "suffix": "",
             "catalog": ("Habitação", "Rendimentos", "Vulnerabilidade")},
    _RSI: {"unit": "pessoas", "dimension": "Total (todos os grupos etários)", "prefix": "", "suffix": "",
           "catalog": ("Proteção Social", "Acesso", "Cuidados")},
    _ESTR: {"unit": "pessoas", "dimension": "Total", "prefix": "", "suffix": "",
            "catalog": ("Proteção Social", "População imigrante", "População imigrante")},
    "Saldo Total": {"unit": "pessoas", "dimension": "Saldo total anual (saldo natural + saldo migratório)", "prefix": "", "suffix": "",
                    "catalog": ("Demografia", "Dinâmica populacional", "Distribuição da população")},
}

INE_DIRECT = {
    "Superfície Km2": {"code": "0011866", "filters": {}},
    "Preço médio m²": {"code": "0012255", "filters": {"dim_3": "H1"}},
    "Rendimento médio": {"code": "0012649", "filters": {"dim_3": "TOT", "dim_4": "T"}},
    "Taxa de analfabetismo": {"code": "0012273", "filters": {"dim_3": "T"}},
    "Proporção de pop. residente em movimentos pendulares": {"code": "0011655", "filters": {}},
}


def normalized(value: object) -> str:
    """Normaliza um texto para comparação: minúsculas, sem acentos nem pontuação e sem espaços repetidos.
    Serve para reconhecer o mesmo território escrito de formas diferentes ("Vila Nova de Cerveira", "VILA NOVA DE CERVEIRA")."""
    text = unicodedata.normalize("NFD", str(value or ""))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", "".join(ch for ch in text if not unicodedata.combining(ch)).lower()).split())


NORMALIZED_TERRITORIES = {normalized(name): name for name in TERRITORIES}
# Antes da NUTS 2013 o INE chamava "Minho-Lima" ao atual Alto Minho (mesmos 10 municípios).
INE_TERRITORIES = {**NORMALIZED_TERRITORIES, "minho lima": "Alto Minho"}


def read_javascript_object(path: Path, variable: str) -> dict:
    """Lê um ficheiro .js do tipo `const variavel = {...};` e devolve o objeto como dicionário Python.
    Dá erro se a variável não existir, para não trabalhar sobre um ficheiro inesperado."""
    source = path.read_text(encoding="utf-8")
    match = re.search(rf"const\s+{re.escape(variable)}\s*=\s*(.*);\s*$", source, re.DOTALL)
    if not match:
        raise RuntimeError(f"Não foi possível ler {variable} em {path.name}")
    return json.loads(match.group(1))


def write_javascript_object(path: Path, comment: str, variable: str, value: dict) -> None:
    """Escreve o dicionário de volta num .js (`// comentário` + `const variavel = {...};`) em UTF-8,
    compacto e sem escapar acentos."""
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    path.write_bytes(f"// {comment}\nconst {variable} = {payload};\n".encode("utf-8"))


def as_number(value: object) -> float | int | None:
    """Converte um valor de célula/JSON em número. Aceita números e textos à portuguesa ("1.234,5", "12 %").
    Devolve inteiro quando não há casas decimais, e None se estiver vazio, for "-", "..", ":" ou não for numérico."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return int(value) if float(value).is_integer() else float(value)
    if isinstance(value, str):
        clean = value.strip().replace(" ", "").replace("%", "").replace("‰", "")
        if not clean or clean in {"-", "..", ":"}:
            return None
        if "," in clean and "." in clean:
            clean = clean.replace(".", "").replace(",", ".")
        elif "," in clean:
            clean = clean.replace(",", ".")
        try:
            number = float(clean)
            return int(number) if number.is_integer() else number
        except ValueError:
            return None
    return None


def validate_values(values: dict[str, float | int], indicator: str) -> None:
    """Confirma que um indicador tem valor numérico para os 11 territórios e que não são todos zero.
    Se falhar, lança ValueError e o indicador não é gravado."""
    missing = [name for name in TERRITORIES if name not in values or as_number(values[name]) is None]
    if missing:
        raise ValueError(f"{indicator}: faltam territórios: {', '.join(missing)}")
    if not any(abs(float(values[name])) > 0 for name in TERRITORIES):
        raise ValueError(f"{indicator}: todos os valores são zero")


def fetch_ine_via_gateway(url: str, attempts: int = 3) -> object:
    """Alternativa para quando o INE bloqueia ligações diretas (servidores partilhados do GitHub).
    Pede o indicador a um gateway de leitura através do curl e só aceita a resposta se ela confirmar o endereço
    oficial do INE e o código exato do indicador. Repete até `attempts` vezes, com espera crescente."""
    code_match = re.search(r"[?&]varcd=(\d+)", url)
    if not code_match:
        raise ValueError(f"Pedido INE sem código de indicador: {url}")
    code = code_match.group(1)
    curl = shutil.which("curl")
    if not curl:
        raise RuntimeError("O cliente curl é necessário para o gateway do INE")
    request_body = json.dumps({
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {"name": "get_indicator", "arguments": {"varcd": code, "lang": "PT"}},
    })
    last_error = None
    for attempt in range(attempts):
        try:
            result = subprocess.run(
                [
                    curl,
                    "--fail",
                    "--silent",
                    "--show-error",
                    "--location",
                    "--compressed",
                    "--connect-timeout",
                    "20",
                    "--max-time",
                    "120",
                    "--header",
                    "Content-Type: application/json",
                    "--data-binary",
                    request_body,
                    "https://gateway.pipeworx.io/ine-pt/mcp",
                ],
                check=True,
                capture_output=True,
                timeout=130,
            )
            envelope = json.loads(result.stdout.decode("utf-8-sig"))
            provenance_url = envelope.get("result", {}).get("_meta", {}).get("provenance", {}).get("upstream", {}).get("url", "")
            if "www.ine.pt/ine/json_indicador/pindica.jsp" not in provenance_url or f"varcd={code}" not in provenance_url:
                raise ValueError("O gateway não confirmou a proveniência oficial do INE")
            content = envelope["result"]["content"][0]["text"]
            payload = json.loads(content)
            if not isinstance(payload, list) or not payload or payload[0].get("IndicadorCod") != code:
                raise ValueError(f"O gateway devolveu um indicador diferente de {code}")
            return payload
        except Exception as error:
            last_error = error
            time.sleep(2**attempt)
    raise RuntimeError(f"Falha ao consultar o indicador INE {code} através do gateway: {last_error}")


def fetch_json_direct(url: str, attempts: int = 3) -> object:
    """Faz um GET direto e devolve o JSON. Usa o curl (IPv4, HTTP/1.1, com User-Agent e Referer do INE)
    e, se o curl não existir, o urllib. Repete até `attempts` vezes, com espera crescente."""
    curl = shutil.which("curl")
    if curl:
        last_error = None
        for attempt in range(attempts):
            try:
                result = subprocess.run(
                    [
                        curl,
                        "--fail",
                        "--silent",
                        "--show-error",
                        "--location",
                        "--ipv4",
                        "--http1.1",
                        "--compressed",
                        "--connect-timeout",
                        "20",
                        "--max-time",
                        "90",
                        "--header",
                        "Accept: application/json, text/javascript, */*; q=0.01",
                        "--header",
                        "Referer: https://www.ine.pt/",
                        "--user-agent",
                        "Mozilla/5.0 (compatible; Observatorio-Alto-Minho/1.0)",
                        url,
                    ],
                    check=True,
                    capture_output=True,
                    timeout=100,
                )
                return json.loads(result.stdout.decode("utf-8-sig"))
            except Exception as error:
                last_error = error
                time.sleep(2**attempt)
        raise RuntimeError(f"Falha ao consultar {url} por IPv4: {last_error}")

    request = urllib.request.Request(url, headers={"User-Agent": "Observatorio-Alto-Minho/1.0"})
    last_error = None
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                return json.load(response)
        except Exception as error:  # a rotina tenta novamente antes de falhar
            last_error = error
            time.sleep(2**attempt)
    raise RuntimeError(f"Falha ao consultar {url}: {last_error}")


def fetch_json(url: str, attempts: int = 3) -> object:
    """Ponto de entrada para obter JSON. No GitHub Actions, os pedidos ao INE tentam primeiro a ligação direta (uma vez)
    e depois o gateway; em todos os outros casos (incluindo no teu computador) a ligação é direta."""
    if os.environ.get("GITHUB_ACTIONS") == "true" and "www.ine.pt/ine/json_indicador/" in url:
        direct_error = None
        try:
            return fetch_json_direct(url, 1)
        except Exception as error:
            direct_error = error
        try:
            return fetch_ine_via_gateway(url, attempts)
        except Exception as gateway_error:
            raise RuntimeError(
                f"Falha na API oficial do INE ({direct_error}) e no gateway alternativo ({gateway_error})"
            ) from gateway_error
    return fetch_json_direct(url, attempts)


def rows_for_latest_complete_year(item: dict, filters: dict[str, str]) -> tuple[str, list[dict]]:
    """Procura, nos dados do INE, o ano mais recente em que existem valores para os 11 territórios,
    depois de aplicar os filtros das dimensões (ex.: {"dim_3": "T"}). Ignora períodos que não sejam um ano de 4 dígitos.
    Devolve (ano, linhas) ou lança ValueError se nenhum ano estiver completo."""
    years = sorted((str(year) for year in item.get("Dados", {}) if re.fullmatch(r"\d{4}", str(year))), reverse=True)
    for year in years:
        rows = [
            row for row in item["Dados"][year]
            if normalized(row.get("geodsg")) in NORMALIZED_TERRITORIES
            and all(str(row.get(key)) == expected for key, expected in filters.items())
            and as_number(row.get("valor")) is not None
        ]
        if len({normalized(row["geodsg"]) for row in rows}) == len(TERRITORIES):
            return year, rows
    raise ValueError("O INE não devolveu um ano completo para os 11 territórios")


def ine_series(item: dict, filters: dict[str, str]) -> dict[str, dict]:
    """Devolve {ano: {território: valor}} para TODOS os anos em que os 11 territórios estão completos."""
    series: dict[str, dict] = {}
    for year, rows in item.get("Dados", {}).items():
        year = str(year)
        if not re.fullmatch(r"\d{4}", year):  # ignora semestres, trimestres, etc.
            continue
        values = {}
        for row in rows:
            territory = INE_TERRITORIES.get(normalized(row.get("geodsg")))
            if territory is None or not all(str(row.get(key)) == expected for key, expected in filters.items()):
                continue
            number = as_number(row.get("valor"))
            if number is not None:
                values[territory] = number
        if len(values) == len(TERRITORIES):
            series[year] = values
    return dict(sorted(series.items()))


def merge_history(entry: dict, new_series: dict[str, dict]) -> int:
    """Junta a série nova ao histórico guardado. Os anos novos substituem os antigos (as fontes revêem valores);
    os anos que a fonte deixou de devolver são mantidos. Só ficam os HISTORY_YEARS anos mais recentes.
    Devolve o número de anos no histórico."""
    merged = {**(entry.get("history") or {}), **new_series}
    entry["history"] = {year: merged[year] for year in sorted(merged)[-HISTORY_YEARS:]}
    return len(entry["history"])


def update_ine(values_by_name: dict, report: list[dict]) -> int:
    """Atualiza os indicadores do INE (os de INE_DIRECT e a diferença salarial H/M) em `values_by_name`.
    Para cada um: grava o ano completo mais recente e o histórico dos anos anteriores. Se o INE só tiver um ano
    mais antigo do que o já guardado, mantém o valor atual. A diferença H/M é calculada a partir dos ganhos de
    homens e mulheres: (H − M) / H × 100. Regista o resultado em `report` e devolve quantos indicadores foram validados."""
    updated = 0
    cache: dict[str, dict] = {}
    for name, config in INE_DIRECT.items():
        code = config["code"]
        if code not in cache:
            payload = fetch_json(f"https://www.ine.pt/ine/json_indicador/pindica.jsp?op=2&varcd={code}&lang=PT")
            if not isinstance(payload, list) or not payload or not payload[0].get("Sucesso", {}).get("Verdadeiro"):
                raise ValueError(f"INE {code}: resposta inválida")
            cache[code] = payload[0]
        year, rows = rows_for_latest_complete_year(cache[code], config["filters"])
        current_year = int(values_by_name[name]["year"])
        if int(year) < current_year:
            report.append({"indicator": name, "source": "INE", "status": "kept", "reason": "source year is older"})
            continue
        new_values = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: as_number(row["valor"]) for row in rows}
        validate_values(new_values, name)
        changed = values_by_name[name]["values"] != new_values or values_by_name[name]["year"] != year
        values_by_name[name]["values"] = new_values
        values_by_name[name]["year"] = year
        history = ine_series(cache[code], config["filters"])
        history[year] = new_values
        years_stored = merge_history(values_by_name[name], history)
        updated += 1
        report.append({"indicator": name, "source": "INE", "status": "updated" if changed else "checked", "year": year, "years_in_history": years_stored})

    name = "Diferença salarial H/M (%)"
    item = cache.get("0012649")
    if item is None:
        payload = fetch_json("https://www.ine.pt/ine/json_indicador/pindica.jsp?op=2&varcd=0012649&lang=PT")
        item = payload[0]
    year_h, men_rows = rows_for_latest_complete_year(item, {"dim_3": "TOT", "dim_4": "1"})
    year_m, women_rows = rows_for_latest_complete_year(item, {"dim_3": "TOT", "dim_4": "2"})
    if year_h != year_m:
        raise ValueError("INE 0012649: os anos por sexo não coincidem")
    men = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: float(row["valor"]) for row in men_rows}
    women = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: float(row["valor"]) for row in women_rows}
    gap = {territory: round((men[territory] - women[territory]) / men[territory] * 100, 4) for territory in TERRITORIES}
    validate_values(gap, name)
    if int(year_h) >= int(values_by_name[name]["year"]):
        changed = values_by_name[name]["values"] != gap or values_by_name[name]["year"] != year_h
        values_by_name[name]["values"] = gap
        values_by_name[name]["year"] = year_h
        men_series = ine_series(item, {"dim_3": "TOT", "dim_4": "1"})
        women_series = ine_series(item, {"dim_3": "TOT", "dim_4": "2"})
        gap_history = {
            y: {t: round((men_series[y][t] - women_series[y][t]) / men_series[y][t] * 100, 4) for t in TERRITORIES}
            for y in men_series
            if y in women_series and all(float(men_series[y][t]) != 0 for t in TERRITORIES)
        }
        gap_history[year_h] = gap
        years_stored = merge_history(values_by_name[name], gap_history)
        updated += 1
        report.append({"indicator": name, "source": "INE", "status": "updated" if changed else "checked", "year": year_h, "years_in_history": years_stored})
    return updated


def read_pordata_layout(path: Path) -> tuple[list, dict[str, int], dict[int, int]]:
    """Lê o Excel da Pordata. Devolve (linhas, {território: linha}, {coluna: ano})."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    territory_rows: dict[str, int] = {}
    for row_index, row in enumerate(rows):
        if len(row) < 2:
            continue
        territory = NORMALIZED_TERRITORIES.get(normalized(row[1]))
        if territory:
            territory_rows.setdefault(territory, row_index)
    if len(territory_rows) != len(TERRITORIES):
        raise ValueError(f"{path.name}: só foram encontrados {len(territory_rows)} territórios")
    year_header_row = next(
        (
            row_index for row_index, row in enumerate(rows)
            if row and normalized(row[0]) == "ambito geografico" and len(row) > 1 and normalized(row[1]) == "anos"
        ),
        None,
    )
    if year_header_row is None:
        raise ValueError(f"{path.name}: cabeçalho de anos não encontrado")
    year_by_column: dict[int, int] = {}
    for column in range(2, max(len(row) for row in rows)):
        year = as_number(rows[year_header_row][column]) if column < len(rows[year_header_row]) else None
        if isinstance(year, int) and 1900 <= year <= 2100:
            year_by_column[column] = year
    return rows, territory_rows, year_by_column


def pordata_candidates(path: Path) -> list[dict]:
    """Lê o Excel da Pordata e devolve uma lista de colunas candidatas, uma por cada coluna de ano
    que tenha valor para os 11 territórios: {"column", "year", "values"}."""
    rows, territory_rows, year_by_column = read_pordata_layout(path)
    candidates = []
    for column, year in year_by_column.items():
        column_values = {}
        for territory, row_index in territory_rows.items():
            cell = rows[row_index][column] if column < len(rows[row_index]) else None
            number = as_number(cell)
            if number is not None:
                column_values[territory] = number
        if len(column_values) == len(TERRITORIES):
            candidates.append({"column": column, "year": year, "values": column_values})
    return candidates


def pordata_history(path: Path, selected_columns: tuple[int, ...], newest_year: int) -> tuple[dict[str, dict], list[int]]:
    """Aplica a mesma escolha de colunas (posição dentro do grupo de colunas de cada ano) a todos os anos.
    Anos com um número diferente de colunas (estrutura alterada) são ignorados em vez de arriscar misturar séries."""
    rows, territory_rows, year_by_column = read_pordata_layout(path)
    groups: dict[int, list[int]] = {}
    for column, year in sorted(year_by_column.items()):
        groups.setdefault(year, []).append(column)
    newest_group = groups[newest_year]
    positions = [newest_group.index(column) for column in selected_columns]
    history: dict[str, dict] = {}
    skipped: list[int] = []
    for year, columns in groups.items():
        if len(columns) != len(newest_group):
            skipped.append(year)
            continue
        picked = [columns[position] for position in positions]
        combined = {}
        for territory, row_index in territory_rows.items():
            parts = [as_number(rows[row_index][c]) if c < len(rows[row_index]) else None for c in picked]
            if any(part is None for part in parts):
                break
            total = round(sum(float(part) for part in parts), 6)
            combined[territory] = int(total) if total.is_integer() else total
        if len(combined) == len(TERRITORIES):
            history[str(year)] = combined
        else:
            skipped.append(year)
    return dict(sorted(history.items())), sorted(skipped)


def candidate_score(candidate: dict[str, float | int], current: dict[str, float | int]) -> float:
    """Mede o quão parecida é uma coluna candidata com os valores atuais: mediana do erro relativo entre
    os 11 territórios. Quanto mais perto de 0, mais provável ser a mesma série."""
    relative_errors = []
    for territory in TERRITORIES:
        old = float(current[territory])
        new = float(candidate[territory])
        relative_errors.append(abs(new - old) / max(abs(old), 1.0))
    return statistics.median(relative_errors)


def choose_pordata_values(path: Path, current: dict, old_year: int) -> tuple[str, dict, float, tuple[int, ...]]:
    """Escolhe no Excel a série que corresponde ao indicador. Um ficheiro pode ter várias colunas para o mesmo ano
    (por exemplo, por sexo ou categoria). Entre as colunas do ano mais recente, experimenta combinações de 1 a 4
    colunas (somadas) e fica com a mais próxima dos valores atuais. Rejeita se o ano for anterior a `old_year`
    ou se o erro mediano passar 0,80. Devolve (ano, valores, erro, colunas escolhidas)."""
    candidates = pordata_candidates(path)
    if not candidates:
        raise ValueError(f"{path.name}: não há colunas de dados completas")
    newest_year = max(candidate["year"] for candidate in candidates)
    if newest_year < old_year:
        raise ValueError(f"{path.name}: o ano {newest_year} é anterior a {old_year}")
    latest = [candidate for candidate in candidates if candidate["year"] == newest_year]
    options = []
    maximum_size = min(4, len(latest))
    for size in range(1, maximum_size + 1):
        for combination in itertools.combinations(latest, size):
            combined = {territory: sum(float(part["values"][territory]) for part in combination) for territory in TERRITORIES}
            combined = {key: int(value) if value.is_integer() else value for key, value in combined.items()}
            options.append((candidate_score(combined, current), combined, tuple(part["column"] for part in combination)))
            if len(options) >= 2000:
                break
        if len(options) >= 2000:
            break
    score, selected, columns = min(options, key=lambda option: option[0])
    if score > 0.80:
        raise ValueError(f"{path.name}: nenhuma série é suficientemente próxima (erro mediano {score:.2f})")
    validate_values(selected, path.name)
    return str(newest_year), selected, score, columns


def choose_fixed_position(path: Path, positions: tuple[int, ...], old_year: int) -> tuple[str, dict, float, tuple[int, ...]]:
    """Soma sempre as colunas nas posições indicadas dentro do grupo de colunas de cada ano (sem procurar a
    combinação mais parecida com os valores antigos). Usa o ano mais recente em que todas essas colunas têm
    valor para os 11 territórios."""
    rows, territory_rows, year_by_column = read_pordata_layout(path)
    groups: dict[int, list[int]] = {}
    for column, year in sorted(year_by_column.items()):
        groups.setdefault(year, []).append(column)
    for year in sorted(groups, reverse=True):
        columns = groups[year]
        if max(positions) >= len(columns):
            continue
        picked = [columns[position] for position in positions]
        values = {}
        for territory, row_index in territory_rows.items():
            parts = [as_number(rows[row_index][c]) if c < len(rows[row_index]) else None for c in picked]
            if all(part is not None for part in parts):
                total = round(sum(float(part) for part in parts), 6)
                values[territory] = int(total) if total.is_integer() else total
        if len(values) == len(TERRITORIES):
            if year < old_year:
                raise ValueError(f"{path.name}: o ano {year} é anterior a {old_year}")
            validate_values(values, path.name)
            return str(year), values, 0.0, tuple(picked)
    raise ValueError(f"{path.name}: não há valores completos para as posições {positions}")


def update_pordata(values_by_name: dict, report: list[dict]) -> int:
    """Atualiza todos os indicadores listados em data/pordata-sources.json a partir dos Excel de .cache/pordata.
    Para cada um: escolhe a série, grava o ano mais recente e o histórico. Se um indicador falhar, mantém os
    valores anteriores e regista o motivo em `report`. No fim, copia a população estrangeira para o seu alias
    ("Pop. Estrangeira c/ estatuto residente"). Devolve quantos indicadores foram validados."""
    sources = json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8"))
    updated = 0
    for name, url in sources.items():
        if name == "Diferença entre o salário mínimo nacional e a remuneração base média mensal dos trabalhadores por conta de outrem":
            continue
        source_id = re.search(r"-(\d+)$", url).group(1)
        path = PORDATA_DIRECTORY / f"{source_id}.xlsx"
        if not path.exists():
            raise FileNotFoundError(f"Falta o ficheiro Pordata {path.name}")
        if name not in values_by_name and name in PORDATA_FIXED_POSITIONS:
            values_by_name[name] = {"year": "0", **{k: v for k, v in PORDATA_FIXED_ENTRIES[name].items() if k != "catalog"}, "values": {}}
        current = values_by_name[name]
        try:
            if name in PORDATA_FIXED_POSITIONS:
                year, selected, score, columns = choose_fixed_position(path, PORDATA_FIXED_POSITIONS[name], int(current["year"]))
            else:
                year, selected, score, columns = choose_pordata_values(path, current["values"], int(current["year"]))
            changed = current["values"] != selected or current["year"] != year
            current["values"] = selected
            current["year"] = year
            history, skipped = pordata_history(path, columns, int(year))
            history[year] = selected
            years_stored = merge_history(current, history)
            updated += 1
            entry_report = {"indicator": name, "source": "Pordata", "status": "updated" if changed else "checked", "year": year, "match_error": round(score, 4), "columns": columns, "years_in_history": years_stored, "years_skipped": skipped}
            if years_stored <= 2:
                entry_report["warning"] = f"só {years_stored} ano(s) no ficheiro {source_id}.xlsx: confirmar se a exportação trouxe todos os anos"
                print(f"[aviso] {name}: {entry_report['warning']}")
            report.append(entry_report)
        except Exception as error:
            report.append({"indicator": name, "source": "Pordata", "status": "kept", "reason": str(error)})
    alias = "Pop. Estrangeira c/ estatuto residente"
    source = "População estrangeira com estatuto legal de residente: total e por algumas nacionalidades"
    values_by_name[alias]["year"] = values_by_name[source]["year"]
    values_by_name[alias]["values"] = dict(values_by_name[source]["values"])
    values_by_name[alias]["history"] = {y: dict(v) for y, v in (values_by_name[source].get("history") or {}).items()}
    return updated


def update_salary_difference(values_by_name: dict, report: list[dict]) -> int:
    """Calcula a diferença entre o rendimento médio (INE) e o salário mínimo nacional (Pordata, ficheiro 74.xlsx),
    para o ano mais recente e para cada ano do histórico em que exista salário mínimo. Lança erro se o salário
    mínimo desse ano não estiver disponível. Devolve 1 quando termina."""
    name = "Diferença entre o salário mínimo nacional e a remuneração base média mensal dos trabalhadores por conta de outrem"
    income = values_by_name["Rendimento médio"]
    year = int(income["year"])
    workbook_path = PORDATA_DIRECTORY / "74.xlsx"
    if not workbook_path.exists():
        raise FileNotFoundError("Falta o ficheiro Pordata 74.xlsx com o salário mínimo")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    wages: dict[int, float | int] = {}
    for row in workbook.worksheets[0].iter_rows(values_only=True):
        if row and len(row) > 2 and isinstance(as_number(row[0]), int) and as_number(row[2]) is not None:
            wages.setdefault(int(as_number(row[0])), as_number(row[2]))
    workbook.close()
    minimum_wage = wages.get(year)
    if minimum_wage is None or minimum_wage <= 0:
        raise ValueError(f"Pordata 74: salário mínimo indisponível para {year}")
    difference = {
        territory: round(float(income["values"][territory]) - float(minimum_wage), 1)
        for territory in TERRITORIES
    }
    validate_values(difference, name)
    changed = values_by_name[name]["values"] != difference or values_by_name[name]["year"] != str(year)
    values_by_name[name]["values"] = difference
    values_by_name[name]["year"] = str(year)
    difference_history = {
        y: {t: round(float(income_values[t]) - float(wages[int(y)]), 1) for t in TERRITORIES}
        for y, income_values in {**(income.get("history") or {}), str(year): income["values"]}.items()
        if int(y) in wages and wages[int(y)] > 0
    }
    merge_history(values_by_name[name], difference_history)
    report.append({
        "indicator": name,
        "source": "Pordata + INE",
        "status": "updated" if changed else "checked",
        "year": str(year),
        "minimum_wage": minimum_wage,
    })
    return 1


def ensure_catalog_entries(values_by_name: dict) -> int:
    """Acrescenta ao catalog-data.js os indicadores novos de PORDATA_FIXED_ENTRIES que ainda não existem
    (só depois de os valores terem passado a validação, para o site nunca mostrar indicadores sem dados).
    Devolve quantos foram acrescentados."""
    catalog = read_javascript_object(CATALOG_PATH, "indicatorCatalog")
    existing = {item["name"] for item in catalog}
    sources = json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8"))
    added = 0
    for name, entry in PORDATA_FIXED_ENTRIES.items():
        if "catalog" not in entry or name in existing or name not in values_by_name:
            continue
        domain, theme, category = entry["catalog"]
        catalog.append({"name": name, "domain": domain, "theme": theme, "category": category,
                        "source": "Pordata", "url": sources[name], "year": values_by_name[name]["year"]})
        added += 1
    if added:
        write_javascript_object(CATALOG_PATH, "Catálogo da matriz intermunicipal e ligações permanentes das fontes.", "indicatorCatalog", catalog)
    return added


def update_catalog_links() -> None:
    """Atualiza no catalog-data.js o endereço de origem de cada indicador: página da Pordata
    (de pordata-sources.json) ou página do indicador no INE."""
    catalog = read_javascript_object(CATALOG_PATH, "indicatorCatalog")
    sources = json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8"))
    for item in catalog:
        if item["name"] in sources:
            item["url"] = sources[item["name"]]
        if item["name"] in INE_DIRECT:
            code = INE_DIRECT[item["name"]]["code"]
            item["url"] = f"https://www.ine.pt/xportal/xmain?xpid=INE&xpgid=ine_indicadores&indOcorrCod={code}&contexto=bd&selTab=tab2"
        if item["name"] == "Diferença salarial H/M (%)":
            item["url"] = "https://www.ine.pt/xportal/xmain?xpid=INE&xpgid=ine_indicadores&indOcorrCod=0012649&contexto=bd&selTab=tab2"
        if item["name"] == "Pop. Estrangeira c/ estatuto residente":
            item["source"] = "Pordata"
            item["url"] = sources["População estrangeira com estatuto legal de residente: total e por algumas nacionalidades"]
    write_javascript_object(CATALOG_PATH, "Catálogo da matriz intermunicipal e ligações permanentes das fontes.", "indicatorCatalog", catalog)


def write_powerbi_dataset(values_by_name: dict) -> None:
    """Escreve powerbi-data.csv: uma linha por indicador e território, só com o ano mais recente.
    É o ficheiro que o relatório do Power BI já usa, por isso o formato não deve mudar."""
    catalog = read_javascript_object(CATALOG_PATH, "indicatorCatalog")
    fields = [
        "Indicador",
        "Territorio",
        "TipoTerritorio",
        "Ano",
        "Valor",
        "Unidade",
        "Dimensao",
        "Dominio",
        "Tema",
        "Categoria",
        "Fonte",
        "FonteURL",
    ]
    with POWERBI_PATH.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for item in catalog:
            dataset = values_by_name[item["name"]]
            for territory in TERRITORIES:
                writer.writerow({
                    "Indicador": item["name"],
                    "Territorio": territory,
                    "TipoTerritorio": "NUTS III" if territory == "Alto Minho" else "Município",
                    "Ano": dataset["year"],
                    "Valor": dataset["values"][territory],
                    "Unidade": dataset["unit"],
                    "Dimensao": dataset["dimension"],
                    "Dominio": item["domain"],
                    "Tema": item["theme"],
                    "Categoria": item["category"],
                    "Fonte": item["source"],
                    "FonteURL": item["url"],
                })


def write_powerbi_history(values_by_name: dict) -> None:
    """Formato longo: uma linha por indicador x território x ano. Ficheiro separado de powerbi-data.csv
    para não duplicar valores nos relatórios Power BI que já usam o ficheiro de «último ano»."""
    catalog = read_javascript_object(CATALOG_PATH, "indicatorCatalog")
    fields = ["Indicador", "Territorio", "TipoTerritorio", "Ano", "Valor", "Unidade", "Dimensao", "Dominio", "Tema", "Categoria", "Fonte", "FonteURL", "AnoMaisRecente"]
    with POWERBI_HISTORY_PATH.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for item in catalog:
            dataset = values_by_name[item["name"]]
            series = {**(dataset.get("history") or {}), dataset["year"]: dataset["values"]}
            for year in sorted(series):
                for territory in TERRITORIES:
                    writer.writerow({
                        "Indicador": item["name"], "Territorio": territory,
                        "TipoTerritorio": "NUTS III" if territory == "Alto Minho" else "Município",
                        "Ano": year, "Valor": series[year][territory], "Unidade": dataset["unit"],
                        "Dimensao": dataset["dimension"], "Dominio": item["domain"], "Tema": item["theme"],
                        "Categoria": item["category"], "Fonte": item["source"], "FonteURL": item["url"],
                        "AnoMaisRecente": "Sim" if year == dataset["year"] else "Nao",
                    })


def main() -> None:
    """Corre a atualização completa. Com `--check-ine` só confirma que o INE responde e termina.
    Caso contrário: atualiza a Pordata e o INE, calcula a diferença para o salário mínimo e escreve o relatório
    (data/update-report.json). Só se todos os indicadores esperados passarem (todos os da Pordata e os 6 do INE)
    é que grava indicator-values.js, catalog-data.js e os dois CSV do Power BI; senão lança erro e deixa os
    ficheiros publicados como estavam."""
    values_by_name = read_javascript_object(VALUES_PATH, "indicatorValuesByName")
    if "--check-ine" in sys.argv:
        code = INE_DIRECT["Superfície Km2"]["code"]
        payload = fetch_json(f"https://www.ine.pt/ine/json_indicador/pindica.jsp?op=2&varcd={code}&lang=PT")
        if not isinstance(payload, list) or not payload or payload[0].get("IndicadorCod") != code:
            raise RuntimeError("A ligação ao INE devolveu uma resposta inválida")
        print("Ligação ao INE e proveniência oficial validadas.")
        return

    expected_pordata = len(json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8")))
    report: list[dict] = []
    pordata_count = update_pordata(values_by_name, report)
    ine_count = update_ine(values_by_name, report)
    pordata_count += update_salary_difference(values_by_name, report)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if pordata_count != expected_pordata:
        failures = [entry for entry in report if entry["source"] == "Pordata" and entry["status"] == "kept"]
        raise RuntimeError(f"Só {pordata_count} de {expected_pordata} indicadores Pordata passaram a validação. Falhas: {failures}")
    if ine_count != 6:
        raise RuntimeError(f"Só {ine_count} indicadores INE passaram a validação")
    ensure_catalog_entries(values_by_name)
    catalog_names = {item["name"] for item in read_javascript_object(CATALOG_PATH, "indicatorCatalog")}
    for stale in [key for key in values_by_name if key not in catalog_names]:
        del values_by_name[stale]
    write_javascript_object(VALUES_PATH, "Valores atualizados automaticamente a partir do INE e da Pordata.", "indicatorValuesByName", values_by_name)
    update_catalog_links()
    write_powerbi_dataset(values_by_name)
    write_powerbi_history(values_by_name)
    print(f"Atualização concluída: {pordata_count} indicadores Pordata e {ine_count} indicadores INE verificados.")


if __name__ == "__main__":
    main()